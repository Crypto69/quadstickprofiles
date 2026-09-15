"""Device-CSV parsing rules that the fixtures cannot show, checked against the firmware:
a mode ends only at an *empty line*; a comma-only row or a row with a blank output is
skipped, not a terminator; anything between the empty line and the next header is
never read. The `.xlsx` path keeps the add-on's rule (stop at the first blank A cell)."""
import pathlib
import pytest
from openpyxl import Workbook
from qsprofile import load, write_csv

FX = pathlib.Path(__file__).parent.parent / "fixtures"

HEAD = (b"QuadStick Configuration,Version 1.4,,Blank rows\r\n"
        b"Profile Name,,Left joy,\r\n"
        b"blank.csv,,Normal,\r\n"
        b"Output or Function,Function,usb,\r\n")


def _write(tmp_path, body):
    src = tmp_path / "blank.csv"
    src.write_bytes(HEAD + body)
    return src


def test_comma_only_row_does_not_end_the_mode(tmp_path):
    src = _write(tmp_path,
                 b"increment_mode,normal,right_sip,\r\n"
                 b",,,,\r\n"
                 b"cross,normal,lip,\r\n"
                 b"\r\n")
    cfg, problems = load(str(src))
    outputs = [m.output for m in cfg.modes[0].mappings]
    assert outputs == ["increment_mode", "cross"]          # cross is still read by the device
    assert not [p for p in problems if p[0] == "error"], problems
    assert any(p[0] == "info" and "no output" in p[3] and "skips it" in p[3] for p in problems)


def test_blank_output_row_with_other_cells_is_skipped_with_a_warning(tmp_path):
    src = _write(tmp_path,
                 b",normal,lip,\r\n"                          # someone forgot the output
                 b"cross,normal,lip,\r\n"
                 b"\r\n")
    cfg, problems = load(str(src))
    assert [m.output for m in cfg.modes[0].mappings] == ["cross"]
    warns = [p for p in problems if p[0] == "warning" and "no output" in p[3]]
    assert warns and warns[0][1] == 1 and warns[0][2] == 4    # mode 1, block row 4
    assert not [p for p in problems if p[0] == "error"]


def test_rows_after_the_empty_line_are_reported_as_ignored(tmp_path):
    src = _write(tmp_path,
                 b"increment_mode,normal,right_sip,\r\n"
                 b"\r\n"
                 b"cross,normal,lip,\r\n"                      # line 7: the device never sees it
                 b"\r\n"
                 b"Profile Name,,Right joy,\r\n"
                 b",,Normal,\r\n"
                 b"Output or Function,Function,usb,\r\n"
                 b"decrement_mode,normal,right_puff,\r\n"
                 b"\r\n")
    cfg, problems = load(str(src))
    assert [m.output for m in cfg.modes[0].mappings] == ["increment_mode"]
    assert [m.output for m in cfg.modes[1].mappings] == ["decrement_mode"]
    errs = [p for p in problems if p[0] == "error"]
    assert len(errs) == 1
    assert "Line 7" in errs[0][3] and "ended mode 1" in errs[0][3] and "ignores it" in errs[0][3]


def test_rows_after_the_preferences_block_are_reported_too(tmp_path):
    src = _write(tmp_path,
                 b"increment_mode,normal,right_sip,\r\n"
                 b"\r\n"
                 b"Preferences,\r\n"
                 b",,,,\r\n"
                 b"Preference,Value,Units,Description,\r\n"
                 b"volume,40,,,\r\n"
                 b"\r\n"
                 b"brightness,9,,,\r\n"
                 b"\r\n")
    cfg, problems = load(str(src))
    assert cfg.preferences == {"volume": "40"}
    errs = [p for p in problems if p[0] == "error"]
    assert len(errs) == 1 and "Preferences block" in errs[0][3]


def test_extra_empty_lines_between_blocks_are_fine(tmp_path):
    src = _write(tmp_path,
                 b"increment_mode,normal,right_sip,\r\n"
                 b"\r\n\r\n\r\n"
                 b"Preferences,\r\n,,,,\r\nPreference,Value,Units,Description,\r\n"
                 b"volume,40,,,\r\n\r\n")
    cfg, problems = load(str(src))
    assert cfg.preferences == {"volume": "40"}
    assert not [p for p in problems if p[0] == "error"], problems


def test_skipped_rows_are_not_written_back(tmp_path):
    """A row the device would skip is dropped on import, so the export is what the
    device actually ran. The mode's other rows are untouched."""
    src = _write(tmp_path,
                 b"increment_mode,normal,right_sip,\r\n"
                 b",,,,\r\n"
                 b"cross,normal,lip,\r\n"
                 b"\r\n"
                 b"Preferences,\r\n,,,,\r\nPreference,Value,Units,Description,\r\n\r\n")
    cfg, _ = load(str(src))
    out = tmp_path / "out.csv"
    write_csv(cfg, str(out))
    assert out.read_bytes() == HEAD + (
        b"increment_mode,normal,right_sip,\r\n"
        b"cross,normal,lip,\r\n"
        b"\r\n"
        b"Preferences,\r\n,,,,\r\nPreference,Value,Units,Description,\r\n\r\n")


def test_xlsx_still_stops_at_the_first_blank_output(tmp_path):
    """The add-on and QMP converters end a sheet at the first blank A cell."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Left joy"
    ws["A1"], ws["C1"] = "Profile Name", "Left joy"
    ws["A2"], ws["C2"] = "blank.csv", "Normal"
    ws["A3"], ws["B3"], ws["C3"] = "Output or Function", "Function", "usb"
    ws.append(["increment_mode", "normal", "right_sip"])
    ws.append([None, None, None])
    ws.append(["cross", "normal", "lip"])
    p = tmp_path / "blank.xlsx"
    wb.save(p)
    cfg, problems = load(str(p))
    assert [m.output for m in cfg.modes[0].mappings] == ["increment_mode"]
    errs = [x for x in problems if x[0] == "error"]
    assert errs and "after the blank row" in errs[0][3]


@pytest.mark.parametrize("name", ["ddfortnite.csv", "cod.csv", "synthetic_pref_override.csv"])
def test_fixtures_have_no_skipped_rows(name):
    """None of the real files carries a blank-output row inside a mode, so the new rule
    changes neither their bytes nor their findings."""
    _, problems = load(str(FX / name))
    assert not [p for p in problems if "no output" in p[3] or "ignores it" in p[3]]


def test_blank_c1_label_stays_blank_through_the_core(tmp_path):
    """`Profile Name,,,` must export as `Profile Name,,,`, never `Profile Name,,Mode 1,`.
    (The API's PUT path is what re-labels it; that fix is in api/, not here.)"""
    src = tmp_path / "nolabel.csv"
    body = (b"QuadStick Configuration,Version 1.4,,No label\r\n"
            b"Profile Name,,,\r\n"
            b"nolabel.csv,,Normal,\r\n"
            b"Output or Function,Function,usb,\r\n"
            b"increment_mode,normal,right_sip,\r\n"
            b"\r\n"
            b"Preferences,\r\n,,,,\r\nPreference,Value,Units,Description,\r\n\r\n")
    src.write_bytes(body)
    cfg, _ = load(str(src))
    assert cfg.modes[0].label == ""
    out = tmp_path / "out.csv"
    write_csv(cfg, str(out))
    assert out.read_bytes() == body


def test_csv_header_only_has_to_start_with_profile(tmp_path):
    """Same firmware rule in the device file: `Profile,,label,` starts a mode."""
    src = tmp_path / "prefix.csv"
    src.write_bytes(b"QuadStick Configuration,Version 1.4,,Prefix\r\n"
                    b"Profile,,Left joy,\r\n"
                    b"prefix.csv,,Normal,\r\n"
                    b"Output or Function,Function,usb,\r\n"
                    b"increment_mode,normal,right_sip,\r\n"
                    b"\r\n"
                    b"Preferences,\r\n,,,,\r\nPreference,Value,Units,Description,\r\nvolume,40,,,\r\n\r\n")
    cfg, problems = load(str(src))
    assert [m.label for m in cfg.modes] == ["Left joy"] and cfg.preferences == {"volume": "40"}
    assert not [p for p in problems if p[0] == "error"]
    infos = [p[3] for p in problems if p[0] == "info" and "Line 2" in p[3]]
    assert infos and "'Profile Name'" in infos[0] and "re-export" in infos[0]
    out = tmp_path / "out.csv"
    write_csv(cfg, str(out))
    assert b"Profile Name,,Left joy," in out.read_bytes()      # canonical on the way out


# ---------------------------------------------------------------- empty function cell (C4)
EMPTY_FN = (b"circle,,lip,\r\n"                            # no function: the device treats it as normal
            b"increment_mode,normal,right_sip,\r\n"
            b"\r\n"
            b"Preferences,\r\n,,,,\r\nPreference,Value,Units,Description,\r\n\r\n")


def test_empty_function_cell_round_trips_byte_for_byte(tmp_path):
    """The parser must not substitute 'normal': the exported file would differ from
    the one on the stick."""
    src = _write(tmp_path, EMPTY_FN)
    cfg, problems = load(str(src))
    m = cfg.modes[0].mappings[0]
    assert (m.output, m.function, m.params, m.inputs) == ("circle", "", [], ["lip"])
    assert not [p for p in problems if p[0] == "error"], problems
    out = tmp_path / "out.csv"
    write_csv(cfg, str(out))
    assert out.read_bytes() == HEAD + EMPTY_FN


def test_empty_function_cell_is_an_info_and_the_card_renders(tmp_path):
    from qsprofile import validate
    from qsprofile.render import render
    cfg, problems = load(str(_write(tmp_path, EMPTY_FN)))
    f = validate(cfg, problems)
    assert not [x for x in f if x[0] == "error"], f
    assert any(x[0] == "info" and x[1] == 1 and x[2] == 4 and "no function" in x[3]
               and "treats it as normal" in x[3] for x in f)
    card = render(cfg, {}, f)
    assert card.startswith("<!doctype html>") and "Row 4 has no function" in card
    assert '<span class="fn"></span>' not in card      # an empty cell gets no function chip


def test_xlsx_keeps_an_empty_function_cell_empty(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Left joy"
    ws["A1"], ws["C1"] = "Profile Name", "Left joy"
    ws["A2"], ws["C2"] = "blank.csv", "Normal"
    ws["A3"], ws["B3"], ws["C3"] = "Output or Function", "Function", "usb"
    ws["A4"], ws["C4"] = "circle", "lip"                     # B4 left empty on purpose
    ws["A5"], ws["B5"], ws["C5"] = "increment_mode", "normal", "right_sip"
    path = tmp_path / "blank.xlsx"
    wb.save(path)
    cfg, _ = load(str(path))
    assert cfg.modes[0].mappings[0].function == ""
    out = tmp_path / "out.csv"
    write_csv(cfg, str(out))
    assert b"circle,,lip,\r\n" in out.read_bytes()


# ---------------------------------------------------------------- per-mode override rows
PREFS_TAIL = b"\r\nPreferences,\r\n,,,,\r\nPreference,Value,Units,Description,\r\n\r\n"


def test_an_output_keyword_is_a_mapping_before_it_is_a_preference_key(tmp_path):
    """The firmware matches column A against its output keywords first, so a name that
    is both an output and a preference key (digital_out_1..4) is an output row, never
    an override. `digital_out_1,,1,` is then a relay with a bogus input `1`, which the
    device ignores; the validator says so as an unknown input, not as an override."""
    from qsprofile import validate
    body = (b"digital_out_1,,1,\r\n"
            b"increment_mode,normal,right_sip,\r\n") + PREFS_TAIL
    cfg, problems = load(str(_write(tmp_path, body)))
    m = cfg.modes[0].mappings[0]
    assert m.kind == "mapping" and m.output == "digital_out_1" and m.inputs == ["1"] and m.value == ""
    out = tmp_path / "out.csv"
    write_csv(cfg, str(out))
    assert out.read_bytes() == HEAD + body
    f = validate(cfg, problems)
    assert [x for x in f if x[0] == "error" and x[2] == 4 and "Unknown input '1'" in x[3]]
    assert not [x for x in f if "for this mode only" in x[3] or "ignores this here" in x[3]]


@pytest.mark.parametrize("n", [1, 2, 3, 4])
def test_every_relay_row_is_a_mapping_and_round_trips(tmp_path, n):
    from qsprofile import validate
    body = (f"digital_out_{n},normal,lip,\r\n".encode()
            + b"increment_mode,normal,right_sip,\r\n") + PREFS_TAIL
    cfg, problems = load(str(_write(tmp_path, body)))
    m = cfg.modes[0].mappings[0]
    assert (m.kind, m.output, m.function, m.inputs) == ("mapping", f"digital_out_{n}", "normal", ["lip"])
    assert not [x for x in validate(cfg, problems) if x[0] == "error"]
    out = tmp_path / "out.csv"
    write_csv(cfg, str(out))
    assert out.read_bytes() == HEAD + body


def test_override_row_keeps_a_filled_column_b_and_warns(tmp_path):
    """The device ignores column B on a preference row and reads the value from C, so
    `mouse_speed,normal,120,` is an override of 120 there. B is kept for byte identity."""
    from qsprofile import validate, write_xlsx
    body = (b"mouse_speed,normal,120,\r\n"
            b"increment_mode,normal,right_sip,\r\n") + PREFS_TAIL
    cfg, problems = load(str(_write(tmp_path, body)))
    m = cfg.modes[0].mappings[0]
    assert (m.kind, m.output, m.function, m.value, m.inputs) == ("preference", "mouse_speed", "normal", "120", [])
    out = tmp_path / "out.csv"
    write_csv(cfg, str(out))
    assert out.read_bytes() == HEAD + body
    f = validate(cfg, problems)
    assert not [x for x in f if x[0] == "error"], f          # not an unknown output any more
    assert [x for x in f if x[0] == "warning" and x[2] == 4 and "column B 'normal'" in x[3]
            and "reads the value from column C" in x[3]]
    xlsx = tmp_path / "out.xlsx"
    write_xlsx(cfg, str(xlsx))
    again, _ = load(str(xlsx))
    r = again.modes[0].mappings[0]
    assert (r.kind, r.function, r.value) == ("preference", "normal", "120")


def test_override_row_with_an_empty_column_b_is_not_warned_about_b(tmp_path):
    from qsprofile import validate
    body = b"mouse_speed,,120,\r\nincrement_mode,normal,right_sip,\r\n" + PREFS_TAIL
    cfg, problems = load(str(_write(tmp_path, body)))
    assert not [x for x in validate(cfg, problems) if "column B" in x[3]]


# ---------------------------------------------------------------- N9: input gaps
def test_gap_between_input_cells_is_closed_up_and_reported(tmp_path):
    """N9: the model stores inputs as a plain ordered list, so a hole between two
    filled cells cannot be kept — and keeping one would break sequence handling.
    The order survives, but the row re-exports shifted left, so say so as an info
    rather than change the file silently."""
    body = (b"increment_mode,normal,right_sip,\r\n"
            b"cross,normal,lip,,right_sip,\r\n"
            + PREFS_TAIL)
    cfg, problems = load(str(_write(tmp_path, body)))
    m = cfg.modes[0].mappings[1]
    assert m.inputs == ["lip", "right_sip"]                  # order kept, gap gone
    gaps = [p for p in problems if "empty input cell" in p[3]]
    assert len(gaps) == 1, problems
    assert gaps[0][0] == "info" and (gaps[0][1], gaps[0][2]) == (1, 5)
    assert "column E" in gaps[0][3]                          # the last filled input column
    assert not [p for p in problems if p[0] == "error"], problems
    out = tmp_path / "out.csv"
    write_csv(cfg, str(out))
    assert b"cross,normal,lip,right_sip,\r\n" in out.read_bytes()


def test_contiguous_inputs_and_trailing_blanks_report_nothing(tmp_path):
    """Only a gap *before* the last filled cell matters; the trailing commas every
    device row carries must never produce a finding."""
    body = (b"increment_mode,normal,right_sip,\r\n"
            b"cross,normal,lip,right_sip,,,\r\n"
            + PREFS_TAIL)
    cfg, problems = load(str(_write(tmp_path, body)))
    assert cfg.modes[0].mappings[1].inputs == ["lip", "right_sip"]
    assert not [p for p in problems if "empty input cell" in p[3]], problems


# ---------------------------------------------------------------- N10: encoding
def test_ansi_saved_csv_is_read_as_windows_1252_with_a_warning(tmp_path):
    """N10: the add-on writes ASCII, but a file re-saved from Excel or Notepad as
    "ANSI" is Windows-1252 and used to raise UnicodeDecodeError before anything
    could be reported. Read it, report where, and let validate() flag the cell —
    write_csv is strict ASCII so the byte can never reach the device anyway."""
    body = b"cross,normal,lip\xe9,\r\nincrement_mode,normal,right_sip,\r\n" + PREFS_TAIL
    cfg, problems = load(str(_write(tmp_path, body)))
    assert cfg.modes[0].mappings[0].inputs == ["lipé"]
    enc = [p for p in problems if "Windows-1252" in p[3]]
    assert len(enc) == 1, problems
    assert enc[0][0] == "warning" and enc[0][2] == 5           # file line, 0xe9 is on line 5
    assert "0xe9" in enc[0][3]
    from qsprofile import validate
    assert [f for f in validate(cfg, problems) if "lipé" in f[3]]


def test_a_utf8_bom_alone_is_not_an_encoding_warning(tmp_path):
    body = b"increment_mode,normal,right_sip,\r\n" + PREFS_TAIL
    src = tmp_path / "bom.csv"
    src.write_bytes(b"\xef\xbb\xbf" + HEAD + body)
    cfg, problems = load(str(src))
    assert cfg.name == "Blank rows"                           # the BOM did not corrupt line 1
    assert not [p for p in problems if "Windows-1252" in p[3]], problems
    assert not [p for p in problems if p[0] == "error"], problems


# ------------------------------------------------------------------- W6: Infrared
def test_an_infrared_block_is_refused_not_parsed_as_a_mode(tmp_path):
    """`write_csv` writes `Profile Name` for every mode, and that header is what the
    firmware dispatches on. Parsing an IR block as a mode would therefore turn it into
    a profile block on export. Count it, error, and store nothing."""
    src = _write(tmp_path,
                 b"cross,normal,lip,\r\n"
                 b"\r\n"
                 b"Infrared,,TV power,\r\n"
                 b"ir_code,,0x20DF10EF,\r\n"
                 b"\r\n")
    cfg, problems = load(str(src))
    assert cfg.infrared_blocks == 1
    assert len(cfg.modes) == 1                      # the IR block is not one of them
    ir = [p for p in problems if "Infrared" in p[3]]
    assert len(ir) == 1 and ir[0][0] == "error", problems
    assert "ir_code" not in [m.output for m in cfg.modes[0].mappings]


def test_a_mode_after_an_infrared_block_keeps_its_numbering(tmp_path):
    """The IR block must not consume a mode number, or every later mode is misreported."""
    src = _write(tmp_path,
                 b"cross,normal,lip,\r\n"
                 b"\r\n"
                 b"Infrared,,TV power,\r\n"
                 b"\r\n"
                 b"Profile Name,,Second,\r\n"
                 b",,Normal,\r\n"
                 b"Output or Function,Function,usb,\r\n"
                 b"circle,normal,lip,\r\n"
                 b"\r\n")
    cfg, problems = load(str(src))
    assert cfg.infrared_blocks == 1
    assert [m.number for m in cfg.modes] == [1, 2]
    assert [m.label for m in cfg.modes] == ["Left joy", "Second"]


def test_a_file_with_no_infrared_block_counts_none(tmp_path):
    cfg, problems = load(str(_write(tmp_path, b"cross,normal,lip,\r\n\r\n")))
    assert cfg.infrared_blocks == 0
    assert not [p for p in problems if "Infrared" in p[3]], problems
