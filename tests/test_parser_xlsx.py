"""Workbook-level parsing rules, checked with small openpyxl workbooks rather than
the fixtures (which are all clean copies of the template)."""
import pytest
from openpyxl import Workbook
from qsprofile import load


def _mode_sheet(ws, title, label, rows):
    ws.title = title
    ws["A1"], ws["C1"] = "Profile Name", label
    ws["A2"], ws["C2"] = "small.csv", "Normal"
    ws["A3"], ws["B3"], ws["C3"] = "Output or Function", "Function", "usb"
    for r in rows:
        ws.append(r)


def _workbook(tmp_path, extra_titles=()):
    wb = Workbook()
    _mode_sheet(wb.active, "Left joy", "Left joy", [["increment_mode", "normal", "right_sip"]])
    for t in extra_titles:
        ws = wb.create_sheet(t)
        ws["A1"], ws["A2"] = "Input keyword", "mp_left_sip"      # a dropdown-list tab
    p = tmp_path / "small.xlsx"
    wb.save(p)
    return p


# ---------------------------------------------------------------- helper tabs
def test_template_helper_tabs_are_skipped_not_errors(tmp_path):
    """A fresh copy of the official template carries Inputs / Outputs / Voice tabs
    holding the dropdown lists. The add-on skips them; an A1 error here would block
    export of every profile started from the template."""
    p = _workbook(tmp_path, ["Inputs", "Outputs", "Voice"])
    cfg, problems = load(str(p))
    assert [m.name for m in cfg.modes] == ["Left joy"]
    assert not [x for x in problems if x[0] == "error"], problems


def test_any_other_unknown_tab_is_still_an_error(tmp_path):
    p = _workbook(tmp_path, ["Inputs", "Scratch"])
    cfg, problems = load(str(p))
    assert [m.name for m in cfg.modes] == ["Left joy"]
    errs = [x for x in problems if x[0] == "error"]
    assert len(errs) == 1 and "'Scratch'" in errs[0][3] and "A1" in errs[0][3]


# ---------------------------------------------------------------- A1 matching
def test_a1_only_has_to_start_with_profile(tmp_path):
    """The firmware dispatches on the start of the line, so `Profile` alone is a mode
    sheet. Say so, and say that re-export writes the canonical `Profile Name`."""
    wb = Workbook()
    _mode_sheet(wb.active, "Left joy", "Left joy", [["increment_mode", "normal", "right_sip"]])
    wb.active["A1"] = "Profile"
    ws = wb.create_sheet("Prefs")
    ws["A1"] = "Preferences of mine"
    ws["A4"], ws["B4"] = "volume", "40"                  # values start at row 4
    p = tmp_path / "prefix.xlsx"
    wb.save(p)
    cfg, problems = load(str(p))
    assert [m.name for m in cfg.modes] == ["Left joy"]
    assert cfg.preferences == {"volume": "40"}
    assert not [x for x in problems if x[0] == "error"], problems
    infos = [x[3] for x in problems if x[0] == "info" and "re-export" in x[3]]
    assert len(infos) == 2
    assert any("'Profile'" in m and "'Profile Name'" in m for m in infos)


@pytest.mark.parametrize("a1", ["profile name", "PROFILE", "Preference", "Infra"])
def test_case_and_shorter_prefixes_are_still_errors(tmp_path, a1):
    wb = Workbook()
    _mode_sheet(wb.active, "Left joy", "Left joy", [["increment_mode", "normal", "right_sip"]])
    wb.active["A1"] = a1
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    cfg, problems = load(str(p))
    assert not cfg.modes
    assert [x for x in problems if x[0] == "error" and "A1" in x[3]]


# ---------------------------------------------------------------- file handle
def _counting_loader(monkeypatch, closed):
    import qsprofile.parser as parser
    real = parser.load_workbook

    def fake(path, **kw):
        wb = real(path, **kw)
        orig = wb.close

        def close():
            closed.append(True)
            orig()
        wb.close = close
        return wb
    monkeypatch.setattr(parser, "load_workbook", fake)


def test_workbook_is_closed_after_parsing(tmp_path, monkeypatch):
    """read_only workbooks hold the zip open until close(); the API parses uploads
    into temp files it then deletes, so a leaked handle is a real leak."""
    closed = []
    _counting_loader(monkeypatch, closed)
    load(str(_workbook(tmp_path)))
    assert closed == [True]


def test_workbook_is_closed_even_when_parsing_fails(tmp_path, monkeypatch):
    import qsprofile.parser as parser
    closed = []
    _counting_loader(monkeypatch, closed)
    monkeypatch.setattr(parser, "hygiene", lambda *a: (_ for _ in ()).throw(RuntimeError("boom")))
    with pytest.raises(RuntimeError):
        load(str(_workbook(tmp_path)))
    assert closed == [True]
