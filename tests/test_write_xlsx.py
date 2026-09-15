"""write_xlsx must never hand openpyxl a sheet title Excel refuses: a mode name is
free text on the device (C1 label / tab name) but a tab name in Excel has rules."""
import pytest
from openpyxl import load_workbook
from qsprofile import load, write_xlsx
from qsprofile.model import Config, Mode, Mapping
from qsprofile.convert import sheet_title


def _cfg(names):
    cfg = Config(name="Titles", filename="titles.csv")
    for i, n in enumerate(names, start=1):
        cfg.modes.append(Mode(number=i, name=n, label=n, channel="usb",
                              mappings=[Mapping(row=4, output="cross", function="normal",
                                                params=[], inputs=["lip"])]))
    return cfg


def _titles(path):
    wb = load_workbook(path, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def test_mode_names_excel_rejects_still_export(tmp_path):
    out = tmp_path / "t.xlsx"
    write_xlsx(_cfg(["A/B?", "[Aim]: fine*", "x" * 40, "Left joy", "left joy", "", "Preferences"]), str(out))
    assert _titles(out) == ["A-B-", "-Aim-- fine-", "x" * 31, "Left joy", "left joy 2",
                            "Mode 6", "Preferences 2", "Preferences"]
    cfg, problems = load(str(out))
    # the C1 label is what the device shows; it is the original text, untouched
    assert [m.label for m in cfg.modes] == ["A/B?", "[Aim]: fine*", "x" * 40, "Left joy", "left joy", "", "Preferences"]
    assert not [p for p in problems if p[0] == "error"], problems


def test_duplicate_suffix_stays_inside_31_characters():
    used = set()
    long = "y" * 31
    assert sheet_title(long, used, "Mode 1") == long
    second = sheet_title(long, used, "Mode 2")
    assert second == "y" * 29 + " 2" and len(second) == 31
    assert sheet_title(long, used, "Mode 3") == "y" * 29 + " 3"


@pytest.mark.parametrize("name,expected", [
    ("Left joy", "Left joy"),
    ("a\\b", "a-b"),
    ("   ", "Mode 9"),
    ("trailing   ", "trailing"),
])
def test_sheet_title_cases(name, expected):
    assert sheet_title(name, set(), "Mode 9") == expected
