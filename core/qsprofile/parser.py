"""Read a QuadStick profile workbook (.xlsx exported from the Google Sheet)
or a device .csv into the model. Layout, from the manual and ddfortnite.xlsx:

  A1 'Profile Name' | C1 mode label      A2 filename (first sheet)  C2 'Normal'
  A3 header row                          C3 usb|bluetooth
  A4.. output | B function | C..J inputs | K.. comments
  The first blank column-A cell ends the sheet.
"""
import csv
from openpyxl import load_workbook
from .model import Config, Mode, Mapping
from .catalog import parse_function, canonical_output, output_label, XBOX_TO_PS, PREFERENCES

SHEET_TYPES = {"Profile Name", "Preferences", "Infrared"}
HELPER_SHEETS = {"inputs", "outputs", "voice"}   # dropdown-list tabs in the official template;
                                                 # the add-on skips them, so does this parser
KEYWORD_COLS = 10          # A..J are read by the device; K onward are comments


def sheet_type(a1):
    """Which block a header starts, or None. The firmware dispatches on the *start* of
    the line, case-sensitive: 'Profile', 'Preferences' or 'Infrared'. Returns the
    canonical header text ('Profile Name', ...); callers report a non-canonical one."""
    for t in SHEET_TYPES:
        if a1.startswith(t.split()[0]):
            return t
    return None


def _noncanonical(where, a1, kind):
    return ("info", None, None,
            f"{where} starts with '{a1}' rather than '{kind}'; the QuadStick accepts anything "
            f"starting with '{kind.split()[0]}', so it is read as a {kind} block. Export writes "
            f"the canonical '{kind}', so the file changes on re-export")


def _cell(v):
    return "" if v is None else str(v).strip()


def _raw(v):
    return "" if v is None else str(v)


def _rows_from_sheet(ws):
    return [[_raw(c) for c in row] for row in ws.iter_rows(values_only=True)]


def hygiene(raw_rows, mode_number, problems):
    """Report what the device would choke on but we repair by stripping: trailing or
    leading whitespace in A..J (the device keeps trailing spaces inside the keyword),
    and newlines inside a cell (the device sees extra lines, and a blank one ends the
    mode early). Returns the stripped rows."""
    clean = []
    for i, row in enumerate(raw_rows):
        for c, v in enumerate(row[:KEYWORD_COLS]):
            if i >= 3 and v and v != v.strip():
                problems.append(("info", mode_number, i + 1,
                                 f"Row {i+1} column {'ABCDEFGHIJ'[c]} had surrounding whitespace in '{v.strip()}'; "
                                 "removed on import (the QuadStick would not have matched it)"))
            if "\n" in v or "\r" in v:
                problems.append(("warning", mode_number, i + 1,
                                 f"Row {i+1} column {'ABCDEFGHIJ'[c]} contains a line break; the QuadStick reads it as "
                                 "extra lines and may stop reading this mode early. Joined into one line on import"))
        clean.append([" ".join(v.split()) if ("\n" in v or "\r" in v) else v.strip() for v in row])
    return clean


def detect_console(rows):
    """A3 says 'XBox Outputs' on Xbox sheets; fall back to sniffing the names."""
    a3 = rows[2][0].lower() if len(rows) > 2 and rows[2] else ""
    if "xbox" in a3:
        return "xbox"
    names = {r[0] for r in rows[3:] if r and r[0]}
    return "xbox" if names & (set(XBOX_TO_PS) - {"start"}) else "playstation"


def parse_mode_rows(rows, number, name, problems, console="playstation", source="xlsx"):
    """rows: list of lists of strings (already stripped). Returns Mode.

    `source` decides what a blank output cell means. In a device `.csv` the firmware
    skips a row with an empty first field and only ends the mode at an empty *line*
    (which `parse_csv` has already used to cut the block), so the row is dropped and
    reported. In an `.xlsx` the add-on and QMP converters stop at the first blank A
    cell and drop everything after it, so that stays the rule there.
    """
    def get(r, c):
        return rows[r][c] if r < len(rows) and c < len(rows[r]) else ""

    mode = Mode(number=number, name=name, label=get(0, 2), channel=get(2, 2).lower())
    if get(0, 10) and get(2, 10).lower() == "usb":
        problems.append(("info", number, 1, "Legacy 'Alternate' block in columns K–R is ignored "
                         "(the QuadStick reads columns K onward as comments)"))
    ended_at = None
    for i in range(3, len(rows)):
        out = get(i, 0)
        if not out:
            if source == "csv":
                filled = any(get(i, c) for c in range(1, KEYWORD_COLS))
                problems.append(("warning" if filled else "info", number, i + 1,
                                 f"Row {i+1} has no output; the QuadStick skips it"
                                 + (" (its other cells are lost)" if filled else "")))
            elif ended_at is None:
                ended_at = i + 1
            continue
        if ended_at is not None:
            problems.append(("error", number, i + 1,
                             f"Row {i+1} has content after the blank row {ended_at}; "
                             "the QuadStick stops reading at the first blank output"))
            continue
        comment = " ".join(get(i, c) for c in range(10, len(rows[i])) if get(i, c))
        output = canonical_output(out, console)
        if output_label(output) is None and out in PREFERENCES:
            # per-mode preference override. The firmware matches column A against its
            # output keywords first, then against its preference names (so digital_out_1..4,
            # which are both, are outputs here, never overrides); on a preference row it
            # ignores column B and reads the value from C. B is kept
            # verbatim in `function` so the file round-trips; validate() warns if it is filled.
            mode.mappings.append(Mapping(row=i + 1, output=out, function=get(i, 1), params=[], inputs=[],
                                         comment=comment, kind="preference", value=get(i, 2)))
            continue
        cell = get(i, 1)
        # an empty B cell is kept empty: the device treats it as normal, and the
        # bytes must round-trip. validate() reports it as an info.
        fname, params = parse_function(cell)[:2] if cell else ("", [])
        inputs = [get(i, c) for c in range(2, 10) if get(i, c)]
        m = Mapping(row=i + 1, output=output, function=fname, params=params, inputs=inputs, comment=comment)
        mode.mappings.append(m)          # a bad function cell is kept as-is; validate() reports it
    return mode


def parse_xlsx(path):
    """Returns (Config, problems) where problems are (severity, mode, row, text)."""
    wb = load_workbook(path, read_only=True)
    try:
        cfg = Config(name=path.rsplit("/", 1)[-1].rsplit(".", 1)[0].replace("_", " "), filename="")
        problems = []
        mode_no = 0
        for ws in wb.worksheets:
            raw = _rows_from_sheet(ws)
            rows = [[_cell(c) for c in r] for r in raw]
            a1 = rows[0][0] if rows and rows[0] else ""
            if ws.title == "Reference Card":
                cfg.reference_card = rows
                continue
            if ws.title.strip().lower() in HELPER_SHEETS:
                continue
            kind = sheet_type(a1)
            if kind is None:
                problems.append(("error", None, 1,
                                 f"Sheet '{ws.title}' cell A1 is '{a1}'; it must be one of "
                                 + ", ".join(sorted(SHEET_TYPES))))
                continue
            if a1 != kind:
                problems.append(_noncanonical(f"Sheet '{ws.title}' cell A1", a1, kind))
            if kind == "Preferences":
                for r in rows[3:]:
                    if r and r[0]:
                        cfg.preferences[r[0]] = r[1] if len(r) > 1 else ""
                continue
            mode_no += 1
            rows = hygiene(raw, mode_no, problems)
            if mode_no == 1:
                cfg.filename = rows[1][0] if len(rows) > 1 and rows[1] else ""
                cfg.console = detect_console(rows)
            cfg.modes.append(parse_mode_rows(rows, mode_no, ws.title, problems, cfg.console))
        return cfg, problems
    finally:
        wb.close()                       # read_only keeps the zip open until told


def parse_csv(path):
    """Device CSV as written by the QuadStick add-on (verified against files
    copied off the flash drive):

      line 1   QuadStick Configuration,Version 1.4,<sheet url>,<sheet name>
      block    Profile Name,,<label>,   /  <filename or blank>,,Normal,  /
               Output or Function,Function,usb,  /  output,function,input,...
      blocks are separated by one blank line; a Preferences block comes last;
      every data line ends with a trailing comma; CRLF line endings; comments
      and the Reference Card sheet are not exported.
    """
    with open(path, newline="", encoding="utf-8-sig") as f:
        allrows = [[_raw(c) for c in r] for r in csv.reader(f)]
    cfg, problems = Config(name=path.rsplit("/", 1)[-1], filename=""), []
    first_line = 1                       # file line number of allrows[0], for messages
    if allrows and allrows[0] and allrows[0][0].strip() == "QuadStick Configuration":
        hdr = allrows[0]
        cfg.name = hdr[3].strip() if len(hdr) > 3 else cfg.name
        cfg.source_url = hdr[2].strip() if len(hdr) > 2 else ""
        cfg.format_version = hdr[1].strip() if len(hdr) > 1 else ""
        allrows = allrows[1:]
        first_line = 2
    # Firmware rules: a block starts at a header line; only an *empty line* (csv.reader
    # gives []) ends it. A comma-only row (',,,,' -> ['', '', ...]) is a row like any
    # other, and lines between the empty line and the next header are never read.
    blocks, cur, ended, modes_seen = [], None, None, 0
    for lineno, r in enumerate(allrows, start=first_line):
        if r == []:
            if cur is not None:
                ended = ("the Preferences block" if cur[0][0].strip().startswith("Preferences")
                         else f"mode {modes_seen}")
            cur = None
            continue
        kind = sheet_type(r[0].strip())
        if kind is not None:
            cur = [r]; blocks.append(cur); ended = None
            if kind != "Preferences":
                modes_seen += 1
            if r[0].strip() != kind:
                problems.append(_noncanonical(f"Line {lineno}", r[0].strip(), kind))
        elif cur is not None:
            cur.append(r)
        elif ended is not None and any(c.strip() for c in r):
            problems.append(("error", None, None,
                             f"Line {lineno} comes after the empty line that ended {ended}; "
                             "the QuadStick ignores it"))
    n = 0
    for raw in blocks:
        if sheet_type(raw[0][0].strip()) == "Preferences":
            for r in raw[2:]:
                r = [_cell(c) for c in r]
                if r and r[0] and r[0] != "Preference":
                    cfg.preferences[r[0]] = r[1] if len(r) > 1 else ""
            continue
        n += 1
        b = hygiene(raw, n, problems)
        if n == 1:
            cfg.filename = b[1][0] if len(b) > 1 and b[1] else ""
            cfg.console = detect_console(b)
        label = b[0][2] if len(b[0]) > 2 else ""
        cfg.modes.append(parse_mode_rows(b, n, label or f"Mode {n}", problems, cfg.console, source="csv"))
    return cfg, problems


def load(path):
    return parse_xlsx(path) if path.lower().endswith(".xlsx") else parse_csv(path)
